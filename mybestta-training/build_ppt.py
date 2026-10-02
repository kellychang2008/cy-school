from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

NAVY=RGBColor(0x1F,0x3A,0x5F); ORANGE=RGBColor(0xF2,0x8C,0x28); LIGHT=RGBColor(0xF5,0xF7,0xFA); GRAY=RGBColor(0x44,0x4B,0x55)
FONT="Microsoft JhengHei"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]

def tb(s,x,y,w,h,text,size=20,color=GRAY,bold=False):
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=t.text_frame; tf.word_wrap=True
    for i,line in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        r=p.add_run(); r.text=line; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name=FONT
        p.space_after=Pt(8)
    return t

def slide(title,body=None,notes=None,dark=False):
    s=prs.slides.add_slide(blank)
    bg=s.background.fill; bg.solid(); bg.fore_color.rgb=NAVY if dark else LIGHT
    if not dark:
        bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(0.25),prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb=ORANGE; bar.line.fill.background()
    tb(s,0.8,0.5,11.5,1.2,title,36,RGBColor(255,255,255) if dark else NAVY,True)
    if body: tb(s,0.8,1.9,11.5,5,body,22,RGBColor(255,255,255) if dark else GRAY)
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s

slide("Mybestta 教務培訓","方便・快速・不失專業\n\n讓老師少做雜事，多教學",dark=True,notes="開場：先問老師們『最近一次被作業或音檔卡到是什麼時候？』")
slide("你是不是也這樣？","・音檔點錯，當場尷尬\n・答案翻半天\n・每晚批改一小時\n・家長問進度，只能憑印象\n・每位老師講法不一樣",notes="讓老師點頭：這些都是日常。")
slide("今天帶走 5 件事","① 上課不翻車\n② 作業秒拍秒改\n③ 依報告帶課\n④ 小考秒拍秒改\n⑤ 家長問不倒＋問官方 Line")
slide("① 上課不翻車","平台使用：開課前 3 分鐘流程\n音檔：從平台入口點，不從資料夾找\n答案：備課標記、現場秒查\n\n鉤子：下一個點錯音檔的，不會是你。",notes="現場示範開課前檢查表。")
slide("防呆口條（團隊一致）","音檔出狀況：「我們先換下一題，這段我會馬上處理。」\n答案不確定：「我先確認標準答案，下課前回覆你。」\n家長問進度：「依平台報告來看，目前……」\n不會用平台：「我立刻請教官方 Line。」",notes="請老師兩兩演練一次。")
slide("② 作業：學生秒拍，老師秒改","學生寫完 → 拍照上傳\n老師集中時段 → 秒改\n回饋三句式：做對什麼＋錯在哪＋下一步\n\n改作業，不該吃掉你的晚餐時間。")
slide("③ 依報告帶課","三個問題：\n1. 哪些題型錯最多？\n2. 哪些學生需要關注？\n3. 哪些已經穩了？\n\n好老師不是更累，是更準。")
slide("20 分鐘檢討流程","0–3 分：公布整體表現\n3–13 分：講錯最多的 2～3 題\n13–18 分：分組／個別補強\n18–20 分：指派下次目標")
slide("④ 小考：秒拍、秒改","小考當天改完，家長當天看到\n短・準・即時\n結果直接串回報告與下一堂開場")
slide("⑤ 家長問不倒","問題統整 → 可回溯 → 回覆有依據\n\n回覆 4 段式：同理 → 依據 → 方向 → 留管道\n\n家長問起，你手上有證據，不是印象。")
slide("卡關怎麼辦？","問官方 Line\n\n不用猜、不用轉來轉去",dark=True)
slide("Before / After","Before：翻找、等待、憑印象、各說各話\n\nAfter：一個入口、即時回饋、依報告、口條一致\n\n[待補：實測節省時間數據]")
slide("你的 7 天上手計畫","Day1 開課前檢查表\nDay2 防呆口條演練\nDay3 作業秒拍秒改上線\nDay4 看第一份報告\nDay5 辦第一次小考\nDay6 練習家長回覆 4 段式\nDay7 回饋給官方 Line")
slide("立刻行動","① 領取 5 本電子書\n② 觀看 3 分鐘教學影片\n③ 加入 Mybestta 官方 Line\n\n方便・快速・不失專業",dark=True)
prs.save("Mybestta_教務培訓.pptx")
